"""유형 변경 뒤 재직 통계의 독립 집계와 미확인 계정의 팀 필터 회귀."""

import io
from contextlib import closing
from urllib.parse import urlencode

import pytest
from openpyxl import load_workbook

from app.services import stats
from tests.factories import make_person, make_team
from tests.test_history_reports import adjustment, moment, record


def money_totals(teams):
    return tuple(
        sum(getattr(team, field) for team in teams)
        for field in ("total_amount", "total_usage", "total_balance")
    )


@pytest.mark.parametrize(
    "monthly_type,monthly_status,balance_type,balance_status",
    [
        ("person", "inactive", "shared", "active"),
        ("shared", "active", "person", "inactive"),
        ("person", "active", "person", "inactive"),
        ("person", "inactive", "person", "active"),
        ("person", "active", "shared", "active"),
        ("shared", "active", "person", "active"),
    ],
)
def test_active_team_financial_totals_follow_each_records_status_and_type(
    auth_client, db, monthly_type, monthly_status, balance_type, balance_status
):
    team = make_team(db, "합성 유형 변경팀")
    person = make_person(
        db, "101", "합성 변경 인원", team=team, account_type=monthly_type, status=monthly_status
    )
    amount = 100 if monthly_status == "active" else 0
    record(db, person, "2026-08", carry=80, amount=amount, usage=20, at=moment(2026, 8, 1))
    person.account_type, person.status = balance_type, balance_status
    db.commit()
    adjustment(db, person, "2026-08", 90, moment(2026, 8, 20))

    for scope in ("observed", "as_of"):
        totals = {}
        for account_type in ("person", "shared", "all"):
            selected_activity = account_type in {"all", monthly_type}
            selected_balance = account_type in {"all", balance_type}
            active_activity = selected_activity and monthly_status == "active"
            active_balance = selected_balance and balance_status == "active"
            expected = (
                amount if active_activity else 0,
                20 if active_activity else 0,
                90 if active_balance else 0,
            )
            report = stats.report(db, "2026-08", scope, account_type)
            totals[account_type] = money_totals(report.active_teams)
            assert totals[account_type] == expected
            assert sum(t.processed_count for t in report.active_teams) == int(active_activity)
            assert sum(t.observed_count for t in report.active_teams) == int(active_balance)
            assert sum(t.count for t in report.active_teams) == int(
                active_activity or active_balance
            )
            # 전체 표와 월간 요약은 재직 여부로 금액을 제외하지 않는다.
            assert money_totals(report.holding_teams) == (
                amount if selected_activity else 0,
                20 if selected_activity else 0,
                90 if selected_balance else 0,
            )
            query = urlencode({"month": "2026-08", "scope": scope, "account_type": account_type})
            page = auth_client.get("/dashboard?" + query)
            assert page.status_code == 200
            assert "충전·순사용은 월간 기록의 재직 상태" in page.text
            exported = auth_client.get("/dashboard/export.xlsx?" + query)
            assert exported.status_code == 200
            with closing(load_workbook(io.BytesIO(exported.content))) as workbook:
                for title, expected_totals in (
                    ("팀별", expected),
                    ("팀별 전체", money_totals(report.holding_teams)),
                ):
                    rows = list(workbook[title].iter_rows(min_row=2, values_only=True))
                    assert (
                        tuple(sum(row[column] or 0 for row in rows) for column in (4, 5, 6))
                        == expected_totals
                    )
        assert totals["all"] == tuple(
            a + b for a, b in zip(totals["person"], totals["shared"], strict=True)
        )
        before_adjustment = stats.report(db, "2026-08", scope, "all", operation_id=0)
        assert money_totals(before_adjustment.active_teams) == (
            (amount, 20, amount + 80) if monthly_status == "active" else (0, 0, 0)
        )
        gap = stats.report(db, "2026-09", scope, "all")
        assert money_totals(gap.active_teams) == (
            0,
            0,
            90 if scope == "as_of" and balance_status == "active" else 0,
        )
        assert money_totals(gap.holding_teams) == (0, 0, 90)


def test_unknown_classification_obeys_current_team_and_person_filters(auth_client, db):
    team_a, team_b = make_team(db, "합성 A팀"), make_team(db, "합성 B팀")
    selected = make_person(db, "101", "합성 A 미확인", team=team_a)
    other = make_person(db, "102", "합성 B 미확인", team=team_b)
    for person, amount in ((selected, 100), (other, 777)):
        record(db, person, "2026-07", amount=amount, profile={}, provenance="unknown")
        record(db, person, "2026-08", amount=amount + 1, profile={}, provenance="unknown")
    for scope in ("observed", "as_of"):
        for account_type in ("person", "shared"):
            for month in ("2026-08", "2026-09"):
                result = stats.report(db, month, scope, account_type, team_name=team_a.name)
                assert [row.person_id for row in result.holding_unclassified_rows] == [selected.id]
                expected_unknown = [selected.id] if scope == "as_of" or month == "2026-08" else []
                assert [row.person_id for row in result.unclassified_rows] == expected_unknown
                query = urlencode(
                    {
                        "month": month,
                        "scope": scope,
                        "account_type": account_type,
                        "team_name": team_a.name,
                    }
                )
                page = auth_client.get("/dashboard?" + query)
                assert page.status_code == 200
                assert "합성 B 미확인" not in page.text
                assert "전체 팀 통계의 계정 유형 미확인 1명" in page.text
                exported = auth_client.get("/dashboard/export.xlsx?" + query)
                with closing(load_workbook(io.BytesIO(exported.content))) as workbook:
                    assert [
                        row[3]
                        for row in workbook["분류 미확인"].iter_rows(min_row=2, values_only=True)
                    ] == ([selected.name] if expected_unknown else [])
            excluded = stats.report(
                db, "2026-08", scope, account_type, person_id=other.id, team_name=team_a.name
            )
            assert excluded.unclassified_rows == excluded.holding_unclassified_rows == []
            missing = stats.report(db, "2026-08", scope, account_type, team_name="없는 팀")
            assert missing.unclassified_rows == missing.holding_unclassified_rows == []
    # 팀 이동 후 미확인 계정도 현재 팀에 따라 조회되며 잔액은 바꾸지 않는다.
    other.team = team_a
    db.commit()
    moved = stats.report(db, "2026-09", team_name=team_a.name)
    assert {row.person_id for row in moved.holding_unclassified_rows} == {selected.id, other.id}
    all_accounts = stats.report(db, "2026-09", account_type="all", team_name=team_a.name)
    assert money_totals(all_accounts.holding_teams) == (0, 0, 879)
