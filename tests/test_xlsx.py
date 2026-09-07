"""표준 XLSX의 실제 파일 입력·보안 경계·보고서 공통 집계 회귀."""

import io
import zipfile
from contextlib import closing

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.models import BalanceRecord, MonthlyDraft
from app.services import stats
from app.services.xlsx import _bytes, read_request, report_workbook, request_template, text_cell
from tests.monthly_helpers import review_fields
from tests.test_ledger_corrections import apply, correction, setup_ledger


def workbook_data(*, month="2026-08", rows=None, modify=None):
    workbook = load_workbook(io.BytesIO(request_template(month)))
    sheet = workbook["요청서"]
    rows = (
        rows
        if rows is not None
        else [[1, "person", "합성팀", "합성이름", "", 100, "0011", "00001101", "비고", 0]]
    )
    for row_number, values in enumerate(rows, 5):
        for column, value in enumerate(values, 1):
            if isinstance(value, str):
                text_cell(sheet, row_number, column, value)
            else:
                sheet.cell(row_number, column, value)
    if modify:
        modify(workbook)
    return _bytes(workbook)


def mutate_zip(data, *, change=None, add=None):
    output = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(data)) as original,
        zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive,
    ):
        for item in original.infolist():
            content = original.read(item)
            if change:
                content = change(item.filename, content)
            archive.writestr(item.filename, content)
        for name, content in (add or {}).items():
            archive.writestr(name, content)
    return output.getvalue()


def test_template_preserves_strings_and_empty_optional_sequence():
    rows = [
        [None, "person", "", "합성", "", 0, "0011", "00000001", "=문자열 메모", None],
        [2, "shared", "", "공용합성", "", 10, None, "00000002", "공용", 0],
    ]
    parsed = read_request(workbook_data(rows=rows), "요청서.xlsx")
    assert parsed.month == "2026-08" and len(parsed.rows) == 2
    assert parsed.rows[0].point_no == "00000001"
    assert parsed.rows[0].personal_no == "0011"
    assert parsed.rows[0].amount == "0" and parsed.rows[0].carry == ""
    assert parsed.rows[0].note == "=문자열 메모"
    assert "H5=00000001" in parsed.rows[0].source_line
    assert parsed.rows[1].personal_no == "" and parsed.rows[1].validated().account_type == "shared"


def test_standard_file_upload_draft_confirmation_roundtrip(auth_client, db):
    rows = [
        [1, "person", "합성팀", "같은이름", "", 100, "0011", "00000001", "일반 비고", 10],
        [2, "person", "합성팀", "같은이름", "", 200, "0011", "00000002", "중복개인번호", 20],
        [3, "shared", "", "공용합성", "", 0, None, "00000003", "공용 비고", 30],
    ]
    response = auth_client.post(
        "/monthly/upload",
        data={"month": "2026-08"},
        files={
            "file": (
                "request.xlsx",
                workbook_data(rows=rows),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.status_code == 200
    values = review_fields(response)
    assert values["input_source"] == "xlsx"
    assert db.get(MonthlyDraft, values["draft_id"]).status == "active"
    assert values["point_no_0"] == "00000001" and values["personal_no_2"] == ""
    values["ack_warnings"] = "yes"
    assert (
        auth_client.post("/monthly/confirm", data=values, follow_redirects=False).status_code == 303
    )
    records = list(db.scalars(select(BalanceRecord).order_by(BalanceRecord.id)))
    assert [record.total for record in records] == [110, 220, 30]
    assert [record.note for record in records] == ["일반 비고", "중복개인번호", "공용 비고"]
    assert stats.report(db, "2026-08", account_type="all").summary.total_balance == 360


def test_multiline_note_form_line_endings_keep_review_and_export(auth_client, db):
    note = "\n=문자열 <비고>\n둘째 줄\n"
    data = workbook_data(rows=[[1, "person", "", "합성", "", 100, "0011", "00001101", note, 0]])
    response = auth_client.post(
        "/monthly/upload", data={"month": "2026-08"}, files={"file": ("test.xlsx", data)}
    )
    assert response.status_code == 200
    values = review_fields(response)
    assert values["note_0"] == note
    # 실제 브라우저는 textarea의 LF를 multipart 전송에서 CRLF로 직렬화한다.
    values["note_0"] = note.replace("\n", "\r\n")
    saved = auth_client.post("/drafts/save", data=values)
    assert saved.status_code == 200
    values.update(saved.json())
    recovered = auth_client.get("/drafts/" + values["draft_id"])
    assert review_fields(recovered)["note_0"] == note
    values["ack_warnings"] = "yes"
    assert (
        auth_client.post("/monthly/confirm", data=values, follow_redirects=False).status_code == 303
    )
    assert db.scalar(select(BalanceRecord.note)) == note
    report = stats.report(db, "2026-08", account_type="all")
    with closing(
        load_workbook(
            io.BytesIO(report_workbook(report, sort_by="name", direction="asc")),
            data_only=False,
        )
    ) as workbook:
        assert workbook["인원"].cell(2, 14).value == note
        assert workbook["인원"].cell(2, 14).data_type == "s"


def test_numeric_identifier_and_formula_keep_cell_reference_until_review(auth_client):
    data = workbook_data(
        rows=[[1, "person", "", "합성", "", 100, "0011", 1101, "", 0]],
        modify=lambda wb: setattr(wb["요청서"]["F5"], "value", "=100+1"),
    )
    result = read_request(data, "test.xlsx")
    assert result.rows[0].point_no == "1101"
    assert result.rows[0].amount == "=100+1"
    assert "F5" in result.rows[0].source_issue and "H5" in result.rows[0].source_issue
    response = auth_client.post(
        "/monthly/upload", data={"month": "2026-08"}, files={"file": ("test.xlsx", data)}
    )
    assert response.status_code == 400
    values = review_fields(response)
    assert values["amount_0"] == "=100+1"
    values.update(amount_0="101", point_no_0="00001101")
    revised = auth_client.post("/monthly/review", data=values)
    assert revised.status_code == 200 and review_fields(revised)["review_token"]


@pytest.mark.parametrize(
    "modify",
    [
        lambda wb: setattr(wb["요청서"]["A1"], "value", "PointBook 월간 보고서"),
        lambda wb: setattr(wb["요청서"]["B1"], "value", "2"),
        lambda wb: setattr(wb["요청서"]["A4"], "value", "다른열"),
        lambda wb: wb.create_sheet("추가"),
        lambda wb: setattr(wb["요청서"]["A2005"], "value", "extra"),
        lambda wb: setattr(wb["요청서"]["K5"], "value", "extra"),
    ],
)
def test_wrong_version_headers_sheets_rows_columns_rejected(modify):
    with pytest.raises(ValueError):
        read_request(workbook_data(modify=modify), "test.xlsx")


@pytest.mark.parametrize("filename", ["request.xlsm", "request.xls", "request.txt"])
def test_non_xlsx_extensions_rejected(filename):
    with pytest.raises(ValueError, match="xlsx"):
        read_request(workbook_data(), filename)


def test_empty_file_partial_row_and_month_mismatch(auth_client):
    with pytest.raises(ValueError, match="인원이 없습니다"):
        read_request(workbook_data(rows=[]), "empty.xlsx")
    parsed = read_request(
        workbook_data(rows=[[None, None, None, None, None, None, None, None, "남은 비고", None]]),
        "partial.xlsx",
    )
    assert len(parsed.rows) == 1 and parsed.rows[0].note == "남은 비고"
    response = auth_client.post(
        "/monthly/upload", data={"month": "2026-09"}, files={"file": ("test.xlsx", workbook_data())}
    )
    assert response.status_code == 400 and "파일 처리 월" in response.text


def test_macro_external_links_and_xml_entities_rejected():
    data = workbook_data()
    with pytest.raises(ValueError, match="매크로"):
        read_request(mutate_zip(data, add={"xl/vbaProject.bin": b"synthetic"}), "test.xlsx")
    external = workbook_data(
        modify=lambda wb: setattr(
            wb["요청서"]["D5"], "hyperlink", "https://example.invalid/synthetic"
        )
    )
    with pytest.raises(ValueError, match="외부 링크"):
        read_request(external, "test.xlsx")

    def entity(name, content):
        if name == "xl/worksheets/sheet1.xml":
            return (
                b'<!DOCTYPE worksheet [<!ENTITY test SYSTEM "file:///synthetic-private">]>'
                + content
            )
        return content

    with pytest.raises(ValueError, match="XML"):
        read_request(mutate_zip(data, change=entity), "test.xlsx")


def test_zip_compressed_unpacked_ratio_and_deadline_limits(monkeypatch):
    from app.services import xlsx

    data = workbook_data()
    monkeypatch.setattr(xlsx, "MAX_ZIP_BYTES", len(data) - 1)
    with pytest.raises(ValueError, match="10MB"):
        read_request(data, "test.xlsx")
    monkeypatch.setattr(xlsx, "MAX_ZIP_BYTES", 10 * 1024 * 1024)
    monkeypatch.setattr(xlsx, "MAX_UNPACKED_BYTES", 100)
    with pytest.raises(ValueError, match="압축 해제"):
        read_request(data, "test.xlsx")
    monkeypatch.setattr(xlsx, "MAX_UNPACKED_BYTES", 30 * 1024 * 1024)
    with pytest.raises(ValueError, match="압축률"):
        read_request(mutate_zip(data, add={"bomb.txt": b"a" * 100000}), "test.xlsx")
    monkeypatch.setattr(xlsx, "MAX_PARSE_SECONDS", -1)
    with pytest.raises(ValueError, match="시간"):
        read_request(data, "test.xlsx")


def test_duplicate_zip_entry_and_malformed_xml_rejected():
    data = workbook_data()
    with pytest.warns(UserWarning):
        duplicated = mutate_zip(data, add={"xl/workbook.xml": b"<root/>"})
    with pytest.raises(ValueError, match="중복"):
        read_request(duplicated, "test.xlsx")
    with pytest.raises(ValueError, match="XML"):
        read_request(
            mutate_zip(
                data, change=lambda name, content: b"<bad" if name == "xl/workbook.xml" else content
            ),
            "test.xlsx",
        )


def test_report_uses_same_scope_cutoff_sort_and_numeric_values(auth_client, db):
    person, _ = setup_ledger(db)
    operation = apply(db, correction(db, person))
    for scope in ["observed", "as_of"]:
        for cutoff in [0, operation.id]:
            result = stats.sort_report(
                stats.report(db, "2026-05", scope, "all", operation_id=cutoff), "total", "desc"
            )
            response = auth_client.get(
                f"/dashboard/export.xlsx?month=2026-05&scope={scope}&account_type=all&operation_id={cutoff}&sort_by=total&direction=desc"
            )
            assert response.status_code == 200 and response.headers["cache-control"] == "no-store"
            with closing(load_workbook(io.BytesIO(response.content), data_only=False)) as wb:
                summary = {row[0].value: row[1].value for row in wb["요약"]}
                assert summary["확인된 잔액 합계"] == result.summary.total_balance
                assert summary["월간 순사용 합계"] == result.summary.total_usage
                assert summary["변경 이력 기준 번호"] == cutoff
                rows = list(wb["인원"].iter_rows(min_row=2, values_only=True))
                assert [row[0] for row in rows] == [row.point_no for row in result.rows]
                assert [row[10] for row in rows] == [row.total for row in result.rows]
            html = auth_client.get(
                f"/dashboard?month=2026-05&scope={scope}&account_type=all&operation_id={cutoff}&sort_by=total&direction=desc"
            )
            assert "export.xlsx?" in html.text


def test_export_untrusted_text_and_large_integer_are_explicit_strings(client, db):
    _person, _ = setup_ledger(db)
    report = stats.report(db, "2026-09", account_type="all")
    report.rows[0].name = '=HYPERLINK("https://example.invalid")'
    report.rows[0].note = "+SUM(1,2)\x01"
    report.summary.total_balance = 1_999_999_999_999_998
    with closing(
        load_workbook(
            io.BytesIO(report_workbook(report, sort_by="name", direction="asc")), data_only=False
        )
    ) as wb:
        assert wb["인원"]["D2"].data_type == "s"
        assert wb["인원"]["O2"].data_type == "s"
        assert "\\u0001" in wb["인원"]["O2"].value
        assert not any(cell.data_type == "f" for sheet in wb for row in sheet for cell in row)
        summary = {row[0].value: row[1] for row in wb["요약"]}
        assert summary["확인된 잔액 합계"].value == "1999999999999998"
        assert summary["확인된 잔액 합계"].data_type == "s"


def test_template_and_export_require_auth_and_validate_parameters(client, auth_client):
    # 별도의 빈 cookie client로 다운로드 인증 확인.
    from fastapi.testclient import TestClient

    with TestClient(auth_client.app) as anonymous:
        assert (
            anonymous.get(
                "/monthly/template.xlsx?month=2026-08", follow_redirects=False
            ).status_code
            == 303
        )
        assert (
            anonymous.get(
                "/dashboard/export.xlsx?month=2026-08", follow_redirects=False
            ).status_code
            == 303
        )
    assert auth_client.get("/monthly/template.xlsx?month=2026-99").status_code == 400
    assert auth_client.get("/dashboard/export.xlsx?month=2026-08&sort_by=bad").status_code == 400
    assert auth_client.get("/dashboard/export.xlsx?month=2026-08&operation_id=x").status_code == 400
    response = auth_client.get("/monthly/template.xlsx?month=2026-08")
    assert response.status_code == 200 and "attachment" in response.headers["content-disposition"]


def test_excel_workers_keep_capacity_until_cancelled_work_finishes(monkeypatch):
    import asyncio
    import threading
    from concurrent.futures import ThreadPoolExecutor

    from app.services import xlsx

    release = threading.Event()
    entered = threading.Barrier(3)
    pool = ThreadPoolExecutor(max_workers=2)
    monkeypatch.setattr(xlsx, "_pool", pool)
    monkeypatch.setattr(xlsx, "_slots", threading.BoundedSemaphore(2))

    def delayed(data, filename):
        entered.wait(timeout=5)
        release.wait(timeout=5)
        return xlsx.XlsxRequest("2026-08", [])

    monkeypatch.setattr(xlsx, "read_request", delayed)

    async def run():
        tasks = [asyncio.create_task(xlsx.extract_request(b"x", "test.xlsx")) for _ in range(2)]
        await asyncio.to_thread(entered.wait, 5)
        tasks[0].cancel()
        with pytest.raises(asyncio.CancelledError):
            await tasks[0]
        with pytest.raises(ValueError, match="다른 Excel"):
            await xlsx.extract_request(b"x", "third.xlsx")
        release.set()
        await tasks[1]

    try:
        asyncio.run(run())
    finally:
        release.set()
        pool.shutdown(wait=True)


def test_explicit_excel_account_type_is_not_silently_overridden(auth_client, db):
    from tests.factories import make_person

    person = make_person(db, "1101", "합성공용", account_type="shared")
    person.personal_no = None
    db.commit()
    response = auth_client.post(
        "/monthly/upload", data={"month": "2026-08"}, files={"file": ("test.xlsx", workbook_data())}
    )
    assert response.status_code == 400
    assert review_fields(response)["account_type_0"] == "person"
    assert "기존 계정 유형과 다릅니다" in response.text
    db.refresh(person)
    assert person.account_type == "shared"
