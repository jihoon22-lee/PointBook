"""두 팀 통계의 모집단·조회 시점·화면/Excel 일치를 합성 장부로 검증한다."""

import io
from contextlib import closing

from openpyxl import load_workbook

from app.services import stats
from app.services.history import preserve_revision
from tests.factories import make_person, make_team
from tests.test_history_reports import operation, record


def team_totals(teams):
    return {
        team.name: (team.count, team.total_amount, team.total_usage, team.total_balance)
        for team in teams
    }


def setup_teams(db):
    team = make_team(db, "합성팀")
    active = make_person(db, "101", "합성 재직", team=team)
    inactive = make_person(db, "102", "합성 비재직", team=team, status="inactive")
    shared = make_person(db, "", "합성 공용", point_no="00000003", account_type="shared")
    unknown = make_person(db, "104", "합성 미관측", team=team)
    unknown.current_amount = 999999
    record(db, active, "2026-07", amount=100)
    record(db, inactive, "2026-07", amount=0, carry=200)
    record(db, shared, "2026-07", amount=300)
    august = record(db, active, "2026-08", amount=50, carry=90, usage=10)
    # 현재 상태·팀 변경이 과거 상태를 덮어쓰면 안 된다.
    active.status = "inactive"
    inactive.status = "active"
    active.team = make_team(db, "현재팀")
    db.commit()
    return active, inactive, shared, unknown, august


def test_team_tables_keep_historical_status_current_team_and_prior_balances(client, db):
    setup_teams(db)
    observed = stats.report(db, "2026-08", account_type="all")
    assert team_totals(observed.active_teams) == {"현재팀": (1, 50, 10, 140)}
    assert team_totals(observed.holding_teams) == {
        "현재팀": (1, 50, 10, 140),
        "합성팀": (2, 0, 0, 200),
        "팀 없음": (1, 0, 0, 300),
    }
    assert observed.summary.total_balance == 140
    assert observed.summary.workforce_count == 1
    assert observed.team_status_unknown_count == 1
    assert next(t for t in observed.holding_teams if t.name == "합성팀").unknown_count == 1
    as_of = stats.report(db, "2026-08", scope="as_of", account_type="all")
    assert team_totals(as_of.active_teams) == {
        "현재팀": (1, 50, 10, 140),
        "팀 없음": (1, 0, 0, 300),
    }
    assert team_totals(as_of.holding_teams) == team_totals(observed.holding_teams)


def test_team_tables_apply_filters_and_same_revision_cutoff(client, db):
    active, inactive, shared, _, august = setup_teams(db)
    op = operation(db, "합성 잔액 정정")
    august.total, august.carry_balance, august.version = 180, 130, 2
    preserve_revision(db, august, op.id, source="manual_correction")
    db.commit()
    for cutoff, total in [(0, 140), (op.id, 180)]:
        result = stats.report(db, "2026-08", person_id=active.id, operation_id=cutoff)
        assert team_totals(result.active_teams) == {"현재팀": (1, 50, 10, total)}
        assert team_totals(result.holding_teams) == team_totals(result.active_teams)
    inactive_only = stats.report(db, "2026-08", person_id=inactive.id)
    assert inactive_only.active_teams == []
    assert team_totals(inactive_only.holding_teams) == {"합성팀": (1, 0, 0, 200)}
    shared_only = stats.report(db, "2026-08", account_type="shared", person_id=shared.id)
    assert team_totals(shared_only.holding_teams) == {"팀 없음": (1, 0, 0, 300)}
    assert team_totals(stats.report(db, "2026-08", team_name="현재팀").holding_teams) == {
        "현재팀": (1, 50, 10, 180)
    }


def test_gap_month_renders_both_team_tables_and_excel_matches(auth_client, db):
    setup_teams(db)
    for month in ["2026-08", "2026-09"]:
        for scope in ["observed", "as_of"]:
            query = f"month={month}&scope={scope}&account_type=all&operation_id=0"
            report = stats.report(db, month, scope, "all", operation_id=0)
            response = auth_client.get("/dashboard?" + query)
            assert response.status_code == 200
            assert 'id="active-team-stats"' in response.text
            assert 'id="all-team-stats"' in response.text
            assert "팀별 재직·비재직 전체 통계" in response.text
            assert "기준시점 재직 상태 미확인 1명" in response.text
            data = auth_client.get("/dashboard/export.xlsx?" + query)
            assert data.status_code == 200
            with closing(load_workbook(io.BytesIO(data.content))) as workbook:
                for sheet, teams in [
                    ("팀별", report.active_teams),
                    ("팀별 전체", report.holding_teams),
                ]:
                    rows = list(workbook[sheet].iter_rows(min_row=2, values_only=True))
                    assert [(r[0], r[1], r[4], r[5], r[6], r[7]) for r in rows] == [
                        (
                            t.name,
                            t.count,
                            t.total_amount if t.processed_count else None,
                            t.total_usage if t.processed_count else None,
                            t.total_balance if t.observed_count else None,
                            t.unknown_count,
                        )
                        for t in teams
                    ]


def test_prior_unclassified_balance_remains_visible_in_team_tables(auth_client, db):
    person = make_person(db, "101", "합성 분류 미확인")
    record(db, person, "2026-07", amount=70, profile={})
    selected = stats.report(db, "2026-08")
    assert selected.active_teams == selected.holding_teams == []
    assert len(selected.holding_unclassified_rows) == 1
    all_accounts = stats.report(db, "2026-08", account_type="all")
    assert team_totals(all_accounts.holding_teams) == {"팀 없음": (1, 0, 0, 70)}
    assert all_accounts.team_status_unknown_count == 1
    response = auth_client.get("/dashboard?month=2026-08")
    assert "전체 팀 통계의 계정 유형 미확인 1명" in response.text
