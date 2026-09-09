"""기존 엑셀 지급 이력·현재 소속·월간 인원 카드의 실제 자료 형태 회귀."""

import io
import json
import re
import uuid
from contextlib import closing

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.models import BalanceRecord, BalanceRevision, MonthlySnapshot
from app.services import stats
from app.services.history import profile_for_person, profile_for_record, record_data
from app.services.observations import latest_observations
from tests.factories import make_person, make_team
from tests.test_ledger_corrections import apply, correction
from tests.test_profiles import form_values


def legacy_record(db, person, month, amount, *, carry=0, usage=0, source="master_at_migration"):
    snapshot = db.scalar(select(MonthlySnapshot).where(MonthlySnapshot.month == month))
    if snapshot is None:
        snapshot = MonthlySnapshot(month=month)
        db.add(snapshot)
        db.flush()
    row = BalanceRecord(
        person_id=person.id,
        snapshot_id=snapshot.id,
        carry_balance=carry,
        amount=amount,
        total=carry + amount,
        usage=usage,
        profile_data=json.dumps(profile_for_person(person)),
        provenance=source,
        version=1,
        note="기존 합성 장부",
    )
    db.add(row)
    db.flush()
    data = record_data(row)
    # 이전 migration처럼 모든 월에 같은 최신 마스터 상태가 들어간 원본 revision을 재현한다.
    data["profile"] = json.loads(row.profile_data)
    db.add(
        BalanceRevision(
            record_id=row.id, version=1, data_json=json.dumps(data), source="migration_baseline"
        )
    )
    db.commit()
    return row


@pytest.mark.parametrize("source", ["master_at_migration", "legacy_import"])
@pytest.mark.parametrize("current_status", ["active", "inactive"])
def test_existing_raw_revisions_restore_absence_employment_departure_and_return(
    auth_client, db, source, current_status
):
    person = make_person(db, "101", "합성 복귀", status=current_status)
    rows = [
        legacy_record(db, person, "2025-01", 0, source=source),
        legacy_record(db, person, "2025-02", 50000, source=source),
        legacy_record(db, person, "2025-03", 0, carry=50000, source=source),
        legacy_record(db, person, "2025-05", 0, carry=50000, source=source),
        legacy_record(db, person, "2025-06", 50000, carry=50000, source=source),
    ]
    originals = [r.profile_data for r in rows]
    revisions = list(db.scalars(select(BalanceRevision.data_json).order_by(BalanceRevision.id)))
    expected = [
        ("inactive", 0, 0),
        ("active", 1, 0),
        ("inactive", 0, 1),
        ("inactive", 0, 0),
        ("active", 1, 0),
    ]
    for row, (status, active, departed) in zip(rows, expected, strict=True):
        assert profile_for_record(row)["status"] == status
        for scope in ("observed", "as_of"):
            result = stats.report(db, row.snapshot.month, scope, operation_id=0)
            assert result.rows[0].status == status
            assert (result.summary.active_count, result.summary.deactivated_count) == (
                active,
                departed,
            )
        assert (
            latest_observations(db, through_month=row.snapshot.month)[person.id].profile["status"]
            == status
        )
    assert [(r.active_count, r.deactivated_count) for r in stats.trend(db)] == [
        (a, d) for _, a, d in expected
    ]
    monthly = auth_client.get("/monthly")
    counts = dict(
        re.findall(
            r'<td class="month-label">(\d{4}-\d{2}).*?</td>\s*<td>(\d+)명</td>',
            monthly.text,
            re.DOTALL,
        )
    )
    assert counts == {
        "2025-01": "0",
        "2025-02": "1",
        "2025-03": "1",
        "2025-05": "0",
        "2025-06": "1",
    }
    assert "<th>총 잔액</th>" in monthly.text and "해당 월 관측 잔액" not in monthly.text
    assert stats.report(db, "2025-04").rows == []
    assert stats.report(db, "2025-04", scope="as_of").summary.active_count == 0
    assert stats.report(db, "2025-04", scope="as_of").summary.total_balance == 50000
    page = auth_client.get(f"/people/{person.id}")
    assert "<th>출처</th>" not in page.text and "이관 시점 인원 정보 참고" not in page.text
    current_table = page.text.split("<h2>현재 정보</h2>")[1].split("</table>")[0]
    assert "잔액 기준 월" not in current_table and "2025-" not in current_table
    assert all(
        label in current_table
        for label in ("<th>이월 잔액</th>", "<th>당월 충전</th>", "<th>총 잔액</th>")
    )
    assert all(
        label not in page.text for label in ("월·정정판", "당시 이름", "당시 팀", "당시 구분")
    )
    values = form_values(auth_client.get(f"/ledger/correct/{person.id}?month=2025-03"))
    assert not any(key.startswith("history_") for key in values)
    assert [r.profile_data for r in rows] == originals
    assert (
        list(db.scalars(select(BalanceRevision.data_json).order_by(BalanceRevision.id)))
        == revisions
    )
    assert person.status == current_status


def test_dashboard_and_xlsx_restore_workforce_counts_and_use_only_current_assignment(
    auth_client, db
):
    old = make_team(db, "과거팀")
    current = make_team(db, "현재팀")
    people = [
        make_person(db, str(i), f"합성{i}", point_no=f"{i:08d}", team=old) for i in range(1, 4)
    ]
    for person, before, now in zip(people, [100, 100, 0], [100, 0, 0], strict=True):
        person.grade = "과거계급"
        legacy_record(db, person, "2025-01", before)
        legacy_record(db, person, "2025-03", now, carry=before)
        person.team, person.grade = current, "현재계급"
    shared = make_person(
        db, "", "합성공용", point_no="00000009", account_type="shared", team=current
    )
    legacy_record(db, shared, "2025-03", 0, carry=200)
    db.commit()
    result = stats.report(db, "2025-03", account_type="all", team_name="현재팀", operation_id=0)
    assert (result.summary.active_count, result.summary.deactivated_count) == (1, 1)
    assert len(result.rows) == 4 and result.summary.total_balance == 500
    assert result.teams[0].name == "현재팀" and result.teams[0].total_balance == 500
    assert not stats.report(db, "2025-03", account_type="all", team_name="과거팀").rows
    assert result.rows[-1].status == "active"  # shared account remains active at zero charge
    monthly = auth_client.get("/monthly")
    counts = dict(
        re.findall(
            r'<td class="month-label">(\d{4}-\d{2}).*?</td>\s*<td>(\d+)명</td>',
            monthly.text,
            re.DOTALL,
        )
    )
    assert counts["2025-03"] == "2"  # active + newly inactive; retained inactive/shared excluded
    html = auth_client.get("/dashboard?month=2025-03&account_type=all")
    assert "재직·비재직 전환 인원" in html.text
    assert "해당 월 관측 잔액" not in html.text
    assert "재직 1명 · 비재직 전환 1명" in html.text
    assert "과거팀" not in html.text and "과거계급" not in html.text
    assert "현재팀" in html.text and "현재계급" in html.text
    data = auth_client.get("/dashboard/export.xlsx?month=2025-03&account_type=all&operation_id=0")
    with closing(load_workbook(io.BytesIO(data.content))) as workbook:
        summary = {row[0].value: row[1].value for row in workbook["요약"]}
        assert (summary["월간 재직 인원"], summary["월간 비재직 전환 인원"]) == (1, 1)
        assert summary["월간 처리 인원"] == 2
        assert summary["잔액 범위"] == "총 잔액"
        team_rows = list(workbook["팀별"].iter_rows(min_row=2, values_only=True))
        assert team_rows[0][2] == 1  # 재직 표에는 신규 비재직 전환 인원을 포함하지 않는다.
        all_team_rows = list(workbook["팀별 전체"].iter_rows(min_row=2, values_only=True))
        assert all_team_rows[0][2] == 2  # 전체 표의 월간 처리 인원은 기존 계약을 유지한다.
        headers = [cell.value for cell in workbook["인원"][1]]
        exported = [
            dict(zip(headers, row, strict=True))
            for row in workbook["인원"].iter_rows(min_row=2, values_only=True)
        ]
        assert {row["현재 팀"] for row in exported} == {"현재팀"}
        assert [r["상태"] for r in exported] == [r.status for r in result.rows]
        assert "당시 정보 출처" not in headers and "월간 당시 팀" not in headers
    page = auth_client.get(f"/ledger/correct/{people[0].id}?month=2025-03")
    assert "history_team_name" not in page.text and "history_grade" not in page.text


def test_corrected_legacy_charge_changes_status_only_in_selected_revision(client, db):
    person = make_person(db, "101", "합성 정정")
    row = legacy_record(db, person, "2025-01", 100)
    next_row = legacy_record(db, person, "2025-03", 0, carry=100)
    person.current_carry_balance, person.current_amount = 100, 0
    db.commit()
    baseline = db.scalar(
        select(BalanceRevision.data_json).where(BalanceRevision.record_id == row.id)
    )
    plan = correction(
        db, person, month="2025-01", carry="100", amount="0", request_key=uuid.uuid4().hex
    )
    assert plan.changes[0]["before"]["profile"]["status"] == "active"
    assert plan.changes[0]["after"]["profile"]["status"] == "inactive"
    op = apply(db, plan)
    assert stats.report(db, "2025-01", operation_id=0).rows[0].status == "active"
    assert stats.report(db, "2025-01", operation_id=op.id).rows[0].status == "inactive"
    assert stats.report(db, "2025-03", operation_id=0).summary.deactivated_count == 1
    assert stats.report(db, "2025-03", operation_id=op.id).summary.deactivated_count == 0
    assert next_row.total == 100 and person.current_carry_balance == 100
    assert (
        db.scalar(
            select(BalanceRevision.data_json).where(
                BalanceRevision.record_id == row.id, BalanceRevision.version == 1
            )
        )
        == baseline
    )
