import io
import json

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.models import BalanceRecord, BalanceRevision, MonthlySnapshot, Person, Team
from app.services.dates import current_month
from app.services.excel_import import import_excel, read_rows
from app.services.stats import report
from app.services.sync import analyze
from app.services.xlsx import report_workbook
from tests.factories import make_person


def _make_xlsx(path, rows, header=None):
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    if header:
        ws.append(header)
    for row in rows:
        ws.append(row)
    wb.save(path)


def test_read_rows_with_headers(tmp_path):
    path = tmp_path / "req.xlsx"
    _make_xlsx(
        path,
        [
            ("1팀", "김소방", "소방경", 50000, "101", "0000 0001", "", 10000),
            ("2팀", "이소방", "소방위", 30000, "102", "0000-0002", "비고", 5000),
        ],
        header=("팀", "이름", "계급", "금액", "개인번호", "포인트번호", "비고", "잔액"),
    )
    rows = read_rows(path)
    assert len(rows) == 2
    assert rows[0].team == "1팀"
    assert rows[0].amount == 50000
    assert rows[0].point_no == "00000001"
    assert rows[0].balance == 10000
    assert rows[1].note == "비고"


def test_read_rows_synonym_headers(tmp_path):
    path = tmp_path / "req.xlsx"
    _make_xlsx(
        path,
        [("구조대", "김소방", "소방경", 50000, "101", "00000001")],
        header=("부서", "성명", "직급", "충전금액", "번호", "포인트 번호"),
    )
    rows = read_rows(path)
    assert rows[0].team == "구조대"
    assert rows[0].personal_no == "101"


def test_read_rows_skips_rows_without_key_fields(tmp_path):
    path = tmp_path / "req.xlsx"
    _make_xlsx(
        path,
        [
            ("1팀", "김소방", "소방경", 50000, "101", "00000001"),
            ("", "", "", "", "", ""),
        ],
        header=("팀", "이름", "계급", "금액", "개인번호", "포인트번호"),
    )
    rows = read_rows(path)
    assert len(rows) == 1


def test_import_excel_creates_teams_persons_snapshot(tmp_path, client, db):
    path = tmp_path / "req.xlsx"
    _make_xlsx(
        path,
        [("1팀", "김소방", "소방경", 50000, "101", "00000001", "", 10000)],
        header=("팀", "이름", "계급", "금액", "개인번호", "포인트번호", "비고", "잔액"),
    )
    result = import_excel(db, path, "2026-07")
    assert result.created_persons == 1
    assert result.records == 1
    person = db.scalar(select(Person).where(Person.point_no == "00000001"))
    assert person is not None
    assert person.team.name == "1팀"
    assert person.status == "active"
    snapshot = db.scalar(select(MonthlySnapshot).where(MonthlySnapshot.month == "2026-07"))
    assert snapshot is not None
    record = snapshot.records[0]
    assert record.carry_balance == 10000
    assert record.total == 60000


def test_import_excel_rejects_nonempty_db(tmp_path, client, db):
    make_person(db, "999", "기존인원")
    path = tmp_path / "req.xlsx"
    _make_xlsx(
        path,
        [("1팀", "김소방", "소방경", 50000, "101")],
        header=("팀", "이름", "계급", "금액", "개인번호"),
    )
    with pytest.raises(ValueError):
        import_excel(db, path, "2026-07")


def test_import_excel_rejects_duplicate_month(tmp_path, client, db):
    path = tmp_path / "req.xlsx"
    _make_xlsx(
        path,
        [("1팀", "김소방", "소방경", 50000, "101", "00000001")],
        header=("팀", "이름", "계급", "금액", "개인번호", "포인트번호"),
    )
    import_excel(db, path, "2026-07")
    with pytest.raises(ValueError):
        import_excel(db, path, "2026-07")


def test_import_excel_team_created_once(tmp_path, client, db):
    path = tmp_path / "req.xlsx"
    _make_xlsx(
        path,
        [
            ("1팀", "김소방", "소방경", 50000, "101", "00000001"),
            ("1팀", "이소방", "소방위", 30000, "102", "00000002"),
        ],
        header=("팀", "이름", "계급", "금액", "개인번호", "포인트번호"),
    )
    import_excel(db, path, "2026-07")
    teams = list(db.scalars(select(Team)).all())
    assert len(teams) == 1


def test_read_rows_rejects_missing_point_number_column(tmp_path):
    path = tmp_path / "missing-point.xlsx"
    _make_xlsx(
        path,
        [("1팀", "김소방", "소방경", 50000, "101")],
        header=("팀", "이름", "계급", "금액", "개인번호"),
    )

    with pytest.raises(ValueError, match="포인트번호"):
        read_rows(path)


@pytest.mark.parametrize("newline", ["\n", "\r\n", "\r"])
def test_import_preserves_note_blank_lines_in_record_revision_and_report(
    tmp_path, client, db, newline
):
    path = tmp_path / "note.xlsx"
    note = newline.join(("", " 첫째 줄 ", "", "둘째 줄", ""))
    expected = "\n 첫째 줄 \n\n둘째 줄\n"
    _make_xlsx(
        path,
        [(" 1팀 ", " 합성이관 ", "소방경", 100, "101", "00000001", note, 20)],
        header=("팀", "이름", "계급", "금액", "개인번호", "포인트번호", "비고", "잔액"),
    )
    import_excel(db, path, "2026-07")
    record = db.scalar(select(BalanceRecord))
    revision = db.scalar(select(BalanceRevision))
    assert record.note == expected
    assert json.loads(revision.data_json)["note"] == expected
    assert record.person.name == "합성이관" and record.person.team.name == "1팀"
    result = report(db, "2026-07")
    assert result.rows[0].activity_note == result.rows[0].note == expected
    workbook = load_workbook(io.BytesIO(report_workbook(result, sort_by="name", direction="asc")))
    sheet = workbook["인원"]
    cells = dict(zip((cell.value for cell in sheet[1]), (cell.value for cell in sheet[2])))
    assert cells["월간 비고"] == cells["잔액 비고"] == expected


def test_import_zero_payment_initializes_inactive_and_preserves_next_month_classification(
    tmp_path, client, db
):
    path = tmp_path / "status.xlsx"
    _make_xlsx(
        path,
        [
            ("1팀", "합성비재직", "소방경", 0, "101", "00000001", "", 100),
            ("1팀", "합성재직", "소방경", 50, "102", "00000002", "", 20),
        ],
        header=("팀", "이름", "계급", "금액", "개인번호", "포인트번호", "비고", "잔액"),
    )
    month = current_month()
    import_excel(db, path, month)
    people = {person.point_no: person for person in db.scalars(select(Person))}
    assert people["00000001"].status == "inactive"
    assert people["00000002"].status == "active"
    assert people["00000001"].current_carry_balance == 100
    result = report(db, month)
    assert result.summary.active_count == result.summary.workforce_count == 1
    assert {row.point_no: row.status for row in result.rows} == {
        point: person.status for point, person in people.items()
    }
    assert {change.point_no: change.action for change in analyze(db, []).changes} == {
        "00000001": "inactive_kept",
        "00000002": "deactivated",
    }
