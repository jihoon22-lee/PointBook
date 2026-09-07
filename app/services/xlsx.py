"""표준 월간 입력과 보고서 XLSX. 외부 실행·수식 평가 없이 한도 안에서 처리한다."""

import asyncio
import io
import time
import zipfile
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass
from datetime import UTC
from threading import BoundedSemaphore
from typing import Any

from defusedxml import ElementTree as SafeET
from defusedxml.common import DefusedXmlException
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string, get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from app._version import __version__
from app.services.dates import KST, validate_month
from app.services.parsing import MAX_REQUEST_ROWS, RawRequestRow
from app.services.stats import Report

FORM_MARKER = "PointBook 월간 요청서"
FORM_VERSION = "1"
HEADERS = (
    "순번",
    "계정 구분",
    "팀",
    "이름",
    "계급",
    "충전액",
    "개인번호",
    "포인트번호",
    "비고",
    "이월 잔액",
)
MAX_ZIP_BYTES = 10 * 1024 * 1024
MAX_UNPACKED_BYTES = 30 * 1024 * 1024
MAX_MEMBER_BYTES = 10 * 1024 * 1024
MAX_PARSE_SECONDS = 8.0
_pool = ThreadPoolExecutor(max_workers=2, thread_name_prefix="pointbook-xlsx")
_slots = BoundedSemaphore(2)


@dataclass
class XlsxRequest:
    month: str
    rows: list[RawRequestRow]


def text_cell(sheet: Worksheet, row: int, column: int, value: object) -> None:
    cell = sheet.cell(row, column)
    cell.value = (
        ""
        if value is None
        else "".join(
            f"\\u{ord(char):04x}" if ord(char) < 32 and char not in "\t\n\r" else char
            for char in str(value)
        )
    )
    cell.data_type = "s"
    cell.number_format = "@"


def _value(sheet: Worksheet, row: int, column: int, value: object) -> None:
    if type(value) is int and abs(value) <= 999_999_999_999_999:
        sheet.cell(row, column, value).number_format = "#,##0;[Red]-#,##0"
    elif value is not None:
        text_cell(sheet, row, column, value)


def _header(sheet: Worksheet, row: int, values: tuple[str, ...]) -> None:
    for column, value in enumerate(values, 1):
        text_cell(sheet, row, column, value)
        cell = sheet.cell(row, column)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="254A73")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(column)].width = 19
    sheet.freeze_panes = f"A{row + 1}"
    sheet.auto_filter.ref = f"A{row}:{get_column_letter(len(values))}{max(row + 1, sheet.max_row)}"


def _bytes(workbook: Workbook) -> bytes:
    stream = io.BytesIO()
    try:
        workbook.save(stream)
        return stream.getvalue()
    finally:
        workbook.close()


def request_template(month: str) -> bytes:
    month = validate_month(month)
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "요청서"
    for row, key, value in [
        (1, FORM_MARKER, FORM_VERSION),
        (2, "처리 월", month),
        (3, "입력 안내", "번호는 텍스트로 보존하고 이월은 모르면 빈칸으로 두세요."),
    ]:
        text_cell(sheet, row, 1, key)
        text_cell(sheet, row, 2, value)
    _header(sheet, 4, HEADERS)
    for row in range(5, 25):
        for column in range(1, 11):
            sheet.cell(row, column).number_format = "@" if column not in {1, 6, 10} else "0"
    choices = DataValidation(type="list", formula1='"person,shared"', allow_blank=False)
    choices.errorTitle, choices.error = "계정 구분", "person 또는 shared를 선택하세요."
    choices.showErrorMessage = True
    sheet.add_data_validation(choices)
    choices.add(f"B5:B{MAX_REQUEST_ROWS + 4}")
    guide = workbook.create_sheet("안내")
    for row, value in enumerate(
        [
            "PointBook 월간 요청서 v1",
            "요청서!B2의 처리 월과 웹의 선택 월을 일치시킵니다.",
            "4행 제목과 1행 양식 버전을 바꾸지 않습니다. 5행부터 최대 2000명입니다.",
            "포인트번호는 구분자를 제외한 8자리 문자열입니다. 선행 0을 보존하세요.",
            "계정 구분은 person(일반), shared(공용). 일반의 개인번호는 필수입니다.",
            "개인번호·이름 중복은 허용하며 공용의 개인번호·팀은 생략할 수 있습니다.",
            "충전액은 0 이상 정수이며 필수입니다. 이월 미입력은 웹 검수에서 채웁니다.",
            "수식·매크로·외부 링크를 사용하지 않습니다. 빈 행 전체만 건너뜁니다.",
            "업로드 뒤 원문 검수·초안·전체 인원 대조를 거쳐 확정합니다.",
            "다운로드 보고서는 이 입력 양식이 아니며 다시 업로드할 수 없습니다.",
        ],
        1,
    ):
        text_cell(guide, row, 1, value)
    guide.column_dimensions["A"].width = 105
    return _bytes(workbook)


def _check_time(started: float) -> None:
    if time.monotonic() - started > MAX_PARSE_SECONDS:
        raise ValueError("Excel 처리 시간이 한도를 넘었습니다. 표준 양식과 행 수를 확인하세요.")


def _archive_check(data: bytes, started: float) -> None:
    if not data or len(data) > MAX_ZIP_BYTES or not data.startswith(b"PK\x03\x04"):
        raise ValueError("표준 .xlsx 파일(최대 10MB)을 선택해 주세요.")
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        items = archive.infolist()
        names = [item.filename for item in items]
        if len(items) > 200 or len(set(names)) != len(names):
            raise ValueError("Excel 압축 항목 수 또는 중복 항목이 허용 범위를 벗어났습니다.")
        if sum(item.file_size for item in items) > MAX_UNPACKED_BYTES:
            raise ValueError("Excel 압축 해제 크기가 30MB 한도를 넘었습니다.")
        sheets = [
            name for name in names if name.startswith("xl/worksheets/") and name.endswith(".xml")
        ]
        if not 1 <= len(sheets) <= 2:
            raise ValueError("표준 요청서와 안내 시트만 사용할 수 있습니다.")
        if any(
            "vbaproject" in name.lower()
            or "externallinks" in name.lower()
            or "activex" in name.lower()
            for name in names
        ):
            raise ValueError("매크로·외부 링크·ActiveX가 포함된 파일은 입력할 수 없습니다.")
        for item in items:
            _check_time(started)
            if item.flag_bits & 1 or item.compress_type not in {
                zipfile.ZIP_STORED,
                zipfile.ZIP_DEFLATED,
            }:
                raise ValueError("암호화 또는 지원하지 않는 압축 형식입니다.")
            if (
                item.file_size > MAX_MEMBER_BYTES
                or item.file_size > max(item.compress_size, 1) * 200
            ):
                raise ValueError("Excel 압축 항목의 크기·압축률 한도를 넘었습니다.")
            if item.filename.endswith((".xml", ".rels")):
                content = archive.read(item)
                root = SafeET.fromstring(content)
                if item.filename.endswith(".rels") and any(
                    element.attrib.get("TargetMode") == "External" for element in root.iter()
                ):
                    raise ValueError("외부 링크가 포함된 Excel은 입력할 수 없습니다.")
                if item.filename == "[Content_Types].xml" and b"macroEnabled" in content:
                    raise ValueError("매크로 형식은 입력할 수 없습니다.")
                if item.filename in sheets:
                    rows = 0
                    last_row = 0
                    coordinates: set[str] = set()
                    for element in root.iter():
                        if element.tag.rsplit("}", 1)[-1] == "row":
                            rows += 1
                            row_number = int(element.attrib.get("r", "0"))
                            if row_number <= last_row:
                                raise ValueError(
                                    "Excel 행 번호가 중복되거나 순서가 잘못되었습니다."
                                )
                            last_row = row_number
                            if (
                                rows > MAX_REQUEST_ROWS + 4
                                or int(element.attrib.get("r", "0")) > MAX_REQUEST_ROWS + 4
                            ):
                                raise ValueError("Excel 행 수 한도(2000명)를 넘었습니다.")
                        elif element.tag.rsplit("}", 1)[-1] == "c":
                            column, row = coordinate_from_string(element.attrib.get("r", "A1"))
                            coordinate = element.attrib.get("r", "")
                            if row != last_row or coordinate in coordinates:
                                raise ValueError(
                                    "Excel 셀 위치가 중복되거나 행과 일치하지 않습니다."
                                )
                            coordinates.add(coordinate)
                            if column_index_from_string(column) > 10 or row > MAX_REQUEST_ROWS + 4:
                                raise ValueError("Excel 열·행이 표준 양식 범위를 벗어났습니다.")


def _raw(value: Any) -> str:
    if value is None:
        return ""
    if type(value) is float and value.is_integer():
        return str(int(value))
    return str(value)


def read_request(data: bytes, filename: str) -> XlsxRequest:
    started = time.monotonic()
    if not filename.lower().endswith(".xlsx"):
        raise ValueError("표준 .xlsx 요청서만 입력할 수 있습니다. .xls/.xlsm은 지원하지 않습니다.")
    workbook = None
    try:
        _archive_check(data, started)
        workbook = load_workbook(
            io.BytesIO(data), read_only=True, data_only=False, keep_links=False
        )
        if set(workbook.sheetnames) not in ({"요청서"}, {"요청서", "안내"}):
            raise ValueError(
                "표준 요청서 시트 이름을 확인하세요. 보고서는 입력 양식으로 사용할 수 없습니다."
            )
        sheet = workbook["요청서"]
        if sheet["A1"].value != FORM_MARKER or str(sheet["B1"].value) != FORM_VERSION:
            raise ValueError(
                "PointBook 월간 요청서 v1 양식이 아닙니다. 표준 템플릿을 내려받으세요."
            )
        month = validate_month(str(sheet["B2"].value or ""))
        if tuple(sheet.cell(4, column).value for column in range(1, 11)) != HEADERS:
            raise ValueError("요청서!4행의 열 제목이 표준 양식과 다릅니다.")
        if sheet.max_row > MAX_REQUEST_ROWS + 4 or sheet.max_column > 10:
            raise ValueError("요청서의 행·열 한도를 넘었습니다.")
        result = []
        for number, cells in enumerate(sheet.iter_rows(min_row=5, max_col=10), 5):
            _check_time(started)
            if all(cell.value is None or cell.value == "" for cell in cells):
                continue
            values = [_raw(cell.value) for cell in cells]
            issues = []
            for cell in cells:
                if cell.data_type == "f":
                    issues.append(
                        f"요청서!{cell.coordinate}: 수식 셀을 실제 값으로 바꾸어 입력하세요."
                    )
            for index in (6, 7):
                if cells[index].value is not None and cells[index].data_type != "s":
                    issues.append(
                        f"요청서!{cells[index].coordinate}: 번호가 문자열이 아닙니다. 선행 0을 추정하지 않았으므로 원본 번호를 확인하세요."
                    )
            account_type = {"일반": "person", "공용": "shared"}.get(values[1], values[1])
            result.append(
                RawRequestRow(
                    account_type=account_type,
                    team=values[2],
                    name=values[3],
                    grade=values[4],
                    amount=values[5],
                    personal_no=values[6],
                    point_no=values[7],
                    note=values[8],
                    carry=values[9],
                    source_line=f"요청서!{number}행 · "
                    + " · ".join(
                        f"{get_column_letter(index)}{number}={value}"
                        for index, value in enumerate(values, 1)
                    ),
                    source_issue=" ".join(issues),
                )
            )
        if not result:
            raise ValueError("요청서에 입력한 인원이 없습니다.")
        return XlsxRequest(month, result)
    except DefusedXmlException as exc:
        raise ValueError("안전하지 않은 XML 선언이 포함된 Excel입니다.") from exc
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(
            "Excel 파일의 압축·XML·셀 구조를 읽을 수 없습니다. 표준 양식을 확인하세요."
        ) from exc
    finally:
        if workbook is not None:
            workbook.close()


async def extract_request(data: bytes, filename: str) -> XlsxRequest:
    if not _slots.acquire(blocking=False):
        raise ValueError("다른 Excel 파일을 처리 중입니다. 잠시 뒤 다시 시도하세요.")
    loop = asyncio.get_running_loop()
    try:
        worker = _pool.submit(read_request, data, filename)
    except Exception:
        _slots.release()
        raise
    worker.add_done_callback(lambda _future: _slots.release())
    wrapped = asyncio.wrap_future(worker, loop=loop)
    wrapped.add_done_callback(lambda done: None if done.cancelled() else done.exception())
    try:
        return await asyncio.wait_for(asyncio.shield(wrapped), timeout=MAX_PARSE_SECONDS + 2)
    except TimeoutError as exc:
        raise ValueError("Excel 처리 시간이 한도를 넘었습니다. 잠시 뒤 다시 시도하세요.") from exc


def _report_workbook(
    report: Report,
    *,
    sort_by: str,
    direction: str,
    team_name: str = "",
    person_id: int | None = None,
) -> bytes:
    if len(report.rows) + len(report.unclassified_rows) > 10000:
        raise ValueError("보고서는 최대 10,000행입니다. 조회 범위를 줄여 주세요.")
    started = time.monotonic()
    workbook = Workbook()
    summary = workbook.active
    assert summary is not None
    summary.title = "요약"
    summary_values = [
        ("PointBook 월간 보고서", __version__),
        ("기준 월", report.month),
        ("잔액 범위", "총 잔액" if report.scope == "observed" else "기준 월까지 마지막 관측"),
        (
            "계정 구분",
            {"person": "일반", "shared": "공용", "all": "일반·공용 전체"}[report.account_type],
        ),
        ("현재 팀 필터", team_name or "전체"),
        ("개인 필터", str(person_id) if person_id else "전체"),
        ("정렬", f"{sort_by}/{direction}"),
        ("변경 이력 기준 번호", report.operation_id or 0),
        (
            "생성 시각 (한국 시간)",
            report.generated_at.replace(tzinfo=UTC).astimezone(KST).isoformat(),
        ),
        ("월간 처리 인원", report.summary.workforce_count),
        ("선택월 잔액 관측 인원", report.summary.observed_count),
        ("확인된 잔액 인원", report.summary.known_balance_count),
        ("기준시점 잔액 미확인", report.summary.unknown_count),
        ("계정 유형 분류 불가", len(report.unclassified_rows)),
        (
            "월간 충전 합계",
            report.summary.total_amount if report.summary.processed_count else "월간 기록 없음",
        ),
        (
            "월간 순사용 합계",
            report.summary.total_usage if report.summary.processed_count else "월간 기록 없음",
        ),
        (
            "확인된 잔액 합계",
            report.summary.total_balance
            if report.summary.known_balance_count
            else "확인된 관측 없음",
        ),
        (
            "범위 안내",
            "팀·계급과 팀별 합계는 현재 인원 정보 기준입니다. 유형은 각 기록을 따르며 미확인 잔액은 합계에서 제외합니다.",
        ),
        (
            "기존 장부 재직 기준",
            "일반 인원은 해당 월 충전액이 0원이면 비재직, 지급액이 있으면 재직입니다. 공용계정은 항상 재직입니다.",
        ),
        ("월간 재직 인원", report.summary.active_count),
        ("월간 비재직 전환 인원", report.summary.deactivated_count),
        ("재직·비재직 전환 인원", report.summary.workforce_count),
        (
            "인원 집계 기준",
            "일반 인원의 해당 월 재직자와 직전 실제 월간 기록에서 재직이었다가 비재직으로 전환한 인원입니다. 기존 비재직·공용계정은 제외합니다.",
        ),
        ("큰 정수", "16자리 이상 정수는 Excel 자릿수 손실을 막기 위해 문자열로 보존합니다."),
        ("용도", "조회 보고서입니다. 월간 요청서 입력 양식으로 사용하지 않습니다."),
    ]
    for summary_row, pair in enumerate(summary_values, 1):
        for column, summary_value in enumerate(pair, 1):
            _value(summary, summary_row, column, summary_value)
    summary.column_dimensions["A"].width, summary.column_dimensions["B"].width = 30, 90
    headers = (
        "포인트번호",
        "계정 구분",
        "개인번호",
        "이름",
        "현재 팀",
        "현재 계급",
        "상태",
        "월간 이월",
        "월간 충전",
        "월간 순사용",
        "선택 범위 잔액",
        "관측 월",
        "관측 구분",
        "월간 비고",
        "잔액 비고",
        "월간 계정 유형",
        "현재 참고 이름",
        "현재 참고 잔액",
    )
    for title, rows in [("인원", report.rows), ("분류 미확인", report.unclassified_rows)]:
        sheet = workbook.create_sheet(title)
        for index, row in enumerate(rows, 2):
            _check_time(started)
            values = (
                row.point_no,
                row.account_type,
                row.personal_no,
                row.name,
                row.team_name,
                row.grade,
                row.status,
                row.carry_balance,
                row.amount,
                row.usage,
                row.total,
                row.observation_month,
                row.balance_kind,
                row.activity_note,
                row.note,
                (row.activity_profile or {}).get("account_type"),
                row.current_reference_name,
                row.current_reference_total,
            )
            for column, cell_value in enumerate(values, 1):
                _value(sheet, index, column, cell_value)
        _header(sheet, 1, headers)
    teams = workbook.create_sheet("팀별")
    for index, team in enumerate(report.teams, 2):
        for column, value in enumerate(
            (
                team.name,
                team.count,
                team.workforce_count,
                team.observed_count,
                team.total_amount if team.processed_count else None,
                team.total_usage if team.processed_count else None,
                team.total_balance if team.observed_count else None,
                team.unknown_count,
            ),
            1,
        ):
            _value(teams, index, column, value)
    _header(
        teams,
        1,
        (
            "현재 팀",
            "조회 인원",
            "월간 처리 인원",
            "잔액 관측 인원",
            "월간 충전",
            "월간 순사용",
            "선택 범위 잔액",
            "미확인 인원",
        ),
    )
    return _bytes(workbook)


def report_workbook(
    report: Report,
    *,
    sort_by: str,
    direction: str,
    team_name: str = "",
    person_id: int | None = None,
) -> bytes:
    if not _slots.acquire(blocking=False):
        raise ValueError("다른 Excel 파일을 처리 중입니다. 잠시 뒤 다시 시도하세요.")
    try:
        return _report_workbook(
            report, sort_by=sort_by, direction=direction, team_name=team_name, person_id=person_id
        )
    finally:
        _slots.release()
